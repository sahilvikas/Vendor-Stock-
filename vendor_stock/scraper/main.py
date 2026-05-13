#!/usr/bin/env python3
"""
Vendor Stock Scraper - Main Entry Point
Run via cron: 30 18 * * * (midnight IST)
Or manually: cd ~/frappe-bench/apps/vendor_stock/vendor_stock/scraper && python3 main.py
"""
import os
import sys
import time
import traceback
import re
from datetime import datetime

# ==========================================
# BOOTSTRAP FRAPPE
# ==========================================
BENCH_PATH = os.path.expanduser("~/frappe-bench")
sys.path.insert(0, os.path.join(BENCH_PATH, "apps", "frappe"))
sys.path.insert(0, os.path.join(BENCH_PATH, "apps", "erpnext"))

import frappe
SITE = "erp.cozycornerpatios.com"
frappe.init(site=SITE, sites_path=os.path.join(BENCH_PATH, "sites"))
frappe.connect()
frappe.set_user("Administrator")

# Now import scrapers (after frappe is initialized)
from vendor_stock.scraper import agora, linen_craft, ddecor, sarom
from vendor_stock.scraper.config import EMAIL_RECIPIENTS, LOW_STOCK_THRESHOLD

# ==========================================
# LOGGING
# ==========================================
SCRAPER_DIR = os.path.dirname(os.path.abspath(__file__))
LOG_DIR = os.path.join(SCRAPER_DIR, "logs")
os.makedirs(LOG_DIR, exist_ok=True)

TIMESTAMP = datetime.now().strftime("%Y%m%d_%H%M")
LOG_FILE = os.path.join(LOG_DIR, f"scrape_{TIMESTAMP}.log")


def log(msg):
    ts = datetime.now().strftime("%H:%M:%S")
    line = f"[{ts}] {msg}"
    print(line, flush=True)
    with open(LOG_FILE, "a") as f:
        f.write(line + "\n")


# ==========================================
# ERP SYNC - Update Item custom fields
# ==========================================
def sync_agora_to_erp(products):
    """Match Agora products to ERP items by code, update vendor stock fields."""
    log("  Syncing Agora to ERP...")
    updated = 0

    # Get all ERP Agora items
    erp_items = frappe.db.sql("""
        SELECT i.name as item_code, i.item_name
        FROM `tabItem` i
        WHERE (i.item_name LIKE 'A - Agora%%' OR i.item_name LIKE 'Agora%%')
        AND i.item_group = 'Raw Material'
        AND i.disabled = 0
    """, as_dict=1)

    # Build code → item_code map
    erp_map = {}
    for item in erp_items:
        m = re.search(r'[Aa]gora[®]?[- ]*(\d+)', item.item_name)
        if m:
            erp_map[m.group(1)] = item.item_code

    now = frappe.utils.now()
    for p in products:
        code = p["code"]
        if code in erp_map:
            try:
                frappe.db.set_value("Item", erp_map[code], {
                    "custom_vendor_stock": p["stock"],
                    "custom_vendor_stock_status": p["status"],
                    "custom_vendor_stock_updated": now,
                    "custom_vendor_stock_price": ""
                }, update_modified=False)
                updated += 1
            except Exception as e:
                log(f"    Error updating {erp_map[code]}: {e}")

    frappe.db.commit()
    log(f"  Agora: {updated}/{len(products)} items updated in ERP")
    return updated


def sync_linen_craft_to_erp(products):
    """Match Linen Craft products to ERP items by keyword, update vendor stock fields."""
    log("  Syncing Linen Craft to ERP...")
    updated = 0

    erp_items = frappe.db.sql("""
        SELECT i.name as item_code, i.item_name
        FROM `tabItem` i
        JOIN `tabItem Default` id ON id.parent = i.name
        WHERE id.default_supplier = 'Linen Craft Pvt Ltd'
        AND i.disabled = 0
    """, as_dict=1)

    now = frappe.utils.now()
    for item in erp_items:
        best_match, best_score = linen_craft.match_erp_to_lc(item.item_name, products)
        if best_match:
            try:
                frappe.db.set_value("Item", item.item_code, {
                    "custom_vendor_stock": best_match["available"],
                    "custom_vendor_stock_status": best_match["status"],
                    "custom_vendor_stock_updated": now,
                    "custom_vendor_stock_price": ""
                }, update_modified=False)
                updated += 1
            except Exception as e:
                log(f"    Error updating {item.item_code}: {e}")

    frappe.db.commit()
    log(f"  Linen Craft: {updated}/{len(erp_items)} items updated in ERP")
    return updated


def sync_ddecor_to_erp(results):
    """Update ERP items with DDécor stock data."""
    log("  Syncing DDécor to ERP...")
    updated = 0
    now = frappe.utils.now()

    for r in results:
        item_code = r.get("item_code", "")
        if not item_code:
            continue

        # check item exists
        if not frappe.db.exists("Item", item_code):
            continue

        stock = 0
        try:
            stock = float(r.get("total_stock", 0))
        except:
            pass

        status = r.get("status", "ERROR")
        price = r.get("price", "")

        try:
            frappe.db.set_value("Item", item_code, {
                "custom_vendor_stock": stock,
                "custom_vendor_stock_status": status,
                "custom_vendor_stock_updated": now,
                "custom_vendor_stock_price": price
            }, update_modified=False)
            updated += 1
        except Exception as e:
            log(f"    Error updating {item_code}: {e}")

    frappe.db.commit()
    log(f"  DDécor: {updated}/{len(results)} items updated in ERP")
    return updated


def sync_sarom_to_erp(matched_items):
    """Update ERP items with Sarom stock data."""
    log("  Syncing Sarom to ERP...")
    updated = 0
    now = frappe.utils.now()

    for m in matched_items:
        item_code = m.get("item_code", "")
        if not item_code:
            continue

        if not frappe.db.exists("Item", item_code):
            continue

        status = m.get("status", "UNKNOWN")

        # Sarom doesn't provide numeric stock, so we use status-based values
        # IN_STOCK = 999 (available), LOW_STOCK = 1 (piece only), OUT_OF_STOCK = 0
        if status == "IN_STOCK":
            stock_val = 999
        elif status == "LOW_STOCK":
            stock_val = 1
        else:
            stock_val = 0

        edd = m.get("edd", "")
        edd_date = None
        if edd:
            try:
                from datetime import datetime as dt
                edd_date = dt.strptime(edd, "%d/%m/%Y").strftime("%Y-%m-%d")
            except:
                pass

        try:
            frappe.db.set_value("Item", item_code, {
                "custom_vendor_stock": stock_val,
                "custom_vendor_stock_status": status,
                "custom_vendor_stock_updated": now,
                "custom_vendor_stock_price": "",
                "custom_vendor_stock_edd": edd_date,
            }, update_modified=False)
            updated += 1
        except Exception as e:
            log(f"    Error updating {item_code}: {e}")

    frappe.db.commit()
    log(f"  Sarom: {updated}/{len(matched_items)} items updated in ERP")
    return updated


# ==========================================
# EXCEL REPORT
# ==========================================
def generate_excel(agora_data, lc_data, ddecor_data, sarom_data):
    """Generate Excel report with all 4 vendors."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    log("  Generating Excel report...")

    GREEN = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    RED = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    YELLOW = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    BOLD = Font(bold=True)

    def style_headers(ws, col_count):
        for c in range(1, col_count + 1):
            ws.cell(row=1, column=c).font = HEADER_FONT
            ws.cell(row=1, column=c).fill = HEADER_FILL

    def auto_width(ws):
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    def color_row(ws, row_num, stock_val, col_count):
        try:
            s = float(stock_val)
            fill = RED if s < LOW_STOCK_THRESHOLD else GREEN
            for c in range(1, col_count + 1):
                ws.cell(row=row_num, column=c).fill = fill
        except:
            pass

    def color_row_sarom(ws, row_num, status, col_count):
        if status == "OUT_OF_STOCK":
            fill = RED
        elif status == "LOW_STOCK":
            fill = YELLOW
        else:
            fill = GREEN
        for c in range(1, col_count + 1):
            ws.cell(row=row_num, column=c).fill = fill

    wb = openpyxl.Workbook()

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Stock Report", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["AGORA"])
    ws["A3"].font = BOLD
    ws.append(["Total", agora_data.get("total", 0)])
    ws.append(["In Stock", agora_data.get("in_stock", 0)])
    ws.append(["Out of Stock", agora_data.get("out_of_stock", 0)])
    ws.append(["Low Stock (<20m)", agora_data.get("low_stock", 0)])
    ws.append([])
    ws.append(["LINEN CRAFT"])
    ws["A9"].font = BOLD
    ws.append(["Total", lc_data.get("total", 0)])
    ws.append(["In Stock", lc_data.get("in_stock", 0)])
    ws.append(["Out of Stock", lc_data.get("out_of_stock", 0)])
    ws.append(["Stock Date", lc_data.get("stock_date", "")])
    ws.append([])
    ws.append(["DDECOR"])
    ws["A15"].font = BOLD
    ws.append(["Total", ddecor_data.get("total", 0)])
    ws.append(["In Stock", ddecor_data.get("in_stock", 0)])
    ws.append(["Out of Stock", ddecor_data.get("out_of_stock", 0)])
    ws.append(["Discontinued", ddecor_data.get("discontinued", 0)])
    ws.append(["Errors", ddecor_data.get("errors", 0)])
    ws.append([])
    ws.append(["SAROM"])
    ws["A22"].font = BOLD
    ws.append(["Total Matched", sarom_data.get("total", 0)])
    ws.append(["In Stock", sarom_data.get("in_stock", 0)])
    ws.append(["Out of Stock", sarom_data.get("out_of_stock", 0)])
    ws.append(["Low Stock (Piece)", sarom_data.get("low_stock", 0)])
    ws.append(["New Email", "Yes" if sarom_data.get("new_email") else "No"])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 15

    # Agora tab
    ws2 = wb.create_sheet("Agora")
    ws2.append(["Code", "Product Name", "Category", "Stock"])
    style_headers(ws2, 4)
    for p in agora_data.get("products", []):
        ws2.append([p["code"], p["name"], p["category"], p["stock"]])
        color_row(ws2, ws2.max_row, p["stock"], 4)
    auto_width(ws2)

    # Linen Craft tab
    ws3 = wb.create_sheet("Linen Craft")
    ws3.append(["Code", "Name", "Width", "Stock", "Committed", "Available"])
    style_headers(ws3, 6)
    for p in lc_data.get("products", []):
        ws3.append([p["code"], p["name"], p["width"], p["stock"], p["committed"], p["available"]])
        color_row(ws3, ws3.max_row, p["available"], 6)
    auto_width(ws3)

    # DDécor tab
    ws4 = wb.create_sheet("DDécor")
    ws4.append(["Item Code", "Item Name", "Collection", "Serial", "Status",
                "Total Stock", "Price", "Error"])
    style_headers(ws4, 8)
    for r in ddecor_data.get("results", []):
        ws4.append([
            r.get("item_code", ""), r.get("item_name", ""),
            r.get("collection_searched", ""), r.get("serial_searched", ""),
            r.get("status", ""), r.get("total_stock", ""),
            r.get("price", ""), r.get("error", "")
        ])
        color_row(ws4, ws4.max_row, r.get("total_stock", ""), 8)
    auto_width(ws4)

    # Sarom tab
    ws5 = wb.create_sheet("Sarom")
    ws5.append(["Item Code", "Item Name", "Collection", "Serial", "Status", "EDD", "Match Method"])
    style_headers(ws5, 7)
    for m in sarom_data.get("matched", []):
        ws5.append([
            m.get("item_code", ""), m.get("item_name", ""),
            m.get("collection", ""), m.get("serial", ""),
            m.get("status", ""), m.get("edd", ""),
            m.get("method", "")
        ])
        color_row_sarom(ws5, ws5.max_row, m.get("status", ""), 7)
    auto_width(ws5)

    excel_path = os.path.join(LOG_DIR, f"stock_report_{TIMESTAMP}.xlsx")
    wb.save(excel_path)
    log(f"  Excel saved: {excel_path}")
    return excel_path


# ==========================================
# SCRAPE LOG - Create/Update DocType record
# ==========================================
def create_scrape_log(triggered_by="Scheduled"):
    """Create a new Vendor Stock Scrape Log doc."""
    doc = frappe.new_doc("Vendor Stock Scrape Log")
    doc.run_date = frappe.utils.today()
    doc.run_start = frappe.utils.now()
    doc.triggered_by = triggered_by
    doc.overall_status = "Running"
    doc.insert(ignore_permissions=True)
    frappe.db.commit()
    log(f"  Created scrape log: {doc.name}")
    return doc


def update_scrape_log(doc, agora_data, lc_data, ddecor_data, sarom_data, erp_updated, excel_path, errors):
    """Update the scrape log with results."""
    doc.run_end = frappe.utils.now()

    # Duration
    try:
        start = frappe.utils.get_datetime(doc.run_start)
        end = frappe.utils.get_datetime(doc.run_end)
        doc.run_duration_mins = round((end - start).total_seconds() / 60, 1)
    except:
        pass

    # Agora
    doc.agora_status = "Success" if agora_data.get("products") else "Failed"
    doc.agora_total = agora_data.get("total", 0)
    doc.agora_in_stock = agora_data.get("in_stock", 0)
    doc.agora_out_of_stock = agora_data.get("out_of_stock", 0)
    doc.agora_low_stock = agora_data.get("low_stock", 0)
    doc.agora_duration_secs = agora_data.get("duration", 0)

    # Linen Craft
    doc.lc_status = "Success" if lc_data.get("products") else "Failed"
    doc.lc_total = lc_data.get("total", 0)
    doc.lc_in_stock = lc_data.get("in_stock", 0)
    doc.lc_out_of_stock = lc_data.get("out_of_stock", 0)
    doc.lc_low_stock = lc_data.get("low_stock", 0)
    doc.lc_stock_date = lc_data.get("stock_date", "")
    doc.lc_duration_secs = lc_data.get("duration", 0)

    # DDécor
    doc.ddecor_status = "Success" if ddecor_data.get("results") else "Failed"
    doc.ddecor_total = ddecor_data.get("total", 0)
    doc.ddecor_in_stock = ddecor_data.get("in_stock", 0)
    doc.ddecor_out_of_stock = ddecor_data.get("out_of_stock", 0)
    doc.ddecor_discontinued = ddecor_data.get("discontinued", 0)
    doc.ddecor_collection_not_found = ddecor_data.get("collection_not_found", 0)
    doc.ddecor_serial_not_found = ddecor_data.get("serial_not_found", 0)
    doc.ddecor_errors = ddecor_data.get("errors", 0)
    doc.ddecor_low_stock = ddecor_data.get("low_stock", 0)
    doc.ddecor_duration_mins = round(ddecor_data.get("duration", 0) / 60, 1)

    # Sarom
    try:
        doc.sarom_status = "Success" if sarom_data.get("matched") else "Failed"
        doc.sarom_total = sarom_data.get("total", 0)
        doc.sarom_in_stock = sarom_data.get("in_stock", 0)
        doc.sarom_out_of_stock = sarom_data.get("out_of_stock", 0)
        doc.sarom_low_stock = sarom_data.get("low_stock", 0)
        doc.sarom_duration_secs = sarom_data.get("duration", 0)
    except Exception:
        # Sarom fields may not exist on the DocType yet - skip gracefully
        pass

    # ERP sync
    doc.erp_items_updated = erp_updated
    doc.erp_sync_status = "Success"

    # Total
    doc.total_items_checked = (agora_data.get("total", 0) +
                                lc_data.get("total", 0) +
                                ddecor_data.get("total", 0) +
                                sarom_data.get("total", 0))

    # Errors
    if errors:
        doc.error_summary = "\n".join(errors)

    # Overall status
    all_ok = (doc.agora_status == "Success" and
              doc.lc_status == "Success" and
              doc.ddecor_status == "Success")
    any_ok = (doc.agora_status == "Success" or
              doc.lc_status == "Success" or
              doc.ddecor_status == "Success")

    if all_ok:
        doc.overall_status = "Success"
    elif any_ok:
        doc.overall_status = "Partial Failure"
    else:
        doc.overall_status = "Failed"

    # Add items to child table
    add_items_to_log(doc, agora_data, lc_data, ddecor_data, sarom_data)

    # Attach Excel
    if excel_path and os.path.exists(excel_path):
        try:
            with open(excel_path, "rb") as f:
                file_doc = frappe.get_doc({
                    "doctype": "File",
                    "file_name": os.path.basename(excel_path),
                    "attached_to_doctype": "Vendor Stock Scrape Log",
                    "attached_to_name": doc.name,
                    "content": f.read(),
                    "is_private": 1
                })
                file_doc.insert(ignore_permissions=True)
                doc.excel_report = file_doc.file_url
        except Exception as e:
            log(f"  Error attaching Excel: {e}")

    doc.save(ignore_permissions=True)
    frappe.db.commit()
    log(f"  Scrape log updated: {doc.name} [{doc.overall_status}]")


def add_items_to_log(doc, agora_data, lc_data, ddecor_data, sarom_data):
    """Add scraped items to the child table with ERP item mapping."""
    import re

    # Build Agora code -> ERP item map
    agora_erp = frappe.db.sql("""
        SELECT i.name as item_code, i.item_name
        FROM `tabItem` i
        WHERE (i.item_name LIKE 'A - Agora%%' OR i.item_name LIKE 'Agora%%')
        AND i.item_group = 'Raw Material'
        AND i.disabled = 0
    """, as_dict=1)
    agora_map = {}
    for item in agora_erp:
        m = re.search(r'[Aa]gora[®]?[- ]*(\d+)', item.item_name)
        if m:
            agora_map[m.group(1)] = {"item_code": item.item_code, "item_name": item.item_name}

    # Build Linen Craft keyword -> ERP item map
    lc_erp = frappe.db.sql("""
        SELECT i.name as item_code, i.item_name
        FROM `tabItem` i
        JOIN `tabItem Default` id ON id.parent = i.name
        WHERE id.default_supplier = 'Linen Craft Pvt Ltd'
        AND i.disabled = 0
    """, as_dict=1)

    # Agora items
    for p in agora_data.get("products", []):
        erp = agora_map.get(p["code"], {})
        doc.append("items", {
            "supplier": "Agora Fabrics",
            "source": "Agora Portal",
            "item_code": erp.get("item_code", ""),
            "item_name": erp.get("item_name", ""),
            "vendor_code": p["code"],
            "vendor_name": p["name"],
            "stock_qty": p["stock"],
            "status": p["status"],
            "erp_updated": 1 if erp.get("item_code") else 0,
            "details": f"Category: {p['category']}"
        })

    # Linen Craft items
    for p in lc_data.get("products", []):
        # match by keyword + fabric code
        best_item = None
        best_score = 0
        for item in lc_erp:
            match, score = linen_craft.match_erp_to_lc(item.item_name, [p])
            if match and score > best_score:
                best_score = score
                best_item = item
        doc.append("items", {
            "supplier": "Linen Craft Pvt Ltd",
            "source": "Linen Craft Sheet",
            "item_code": best_item.item_code if best_item and best_score >= 1 else "",
            "item_name": best_item.item_name if best_item and best_score >= 1 else "",
            "vendor_code": p["code"],
            "vendor_name": p["name"],
            "stock_qty": p["available"],
            "status": p["status"],
            "erp_updated": 1 if best_item and best_score >= 1 else 0,
            "details": f"Stock: {p['stock']}, Committed: {p['committed']}, Available: {p['available']}"
        })

    # DDécor items
    for r in ddecor_data.get("results", []):
        stock = 0
        try:
            stock = float(r.get("total_stock", 0))
        except:
            pass
        doc.append("items", {
            "supplier": r.get("supplier", ""),
            "source": "DDécor Portal",
            "item_code": r.get("item_code", ""),
            "item_name": r.get("item_name", ""),
            "vendor_code": f"{r.get('collection_searched', '')} / {r.get('serial_searched', '')}",
            "stock_qty": stock,
            "price": r.get("price", ""),
            "status": r.get("status", ""),
            "erp_updated": 1,
            "error": r.get("error", ""),
            "details": f"Matched: {r.get('product_name', '')}"
        })

    # Sarom items
    for m in sarom_data.get("matched", []):
        stock_val = 999 if m["status"] == "IN_STOCK" else (1 if m["status"] == "LOW_STOCK" else 0)
        doc.append("items", {
            "supplier": "Sarom",
            "source": "Sarom Email",
            "item_code": m.get("item_code", ""),
            "item_name": m.get("item_name", ""),
            "vendor_code": f"{m.get('collection', '')} / {m.get('serial', '')}",
            "stock_qty": stock_val,
            "status": m.get("status", ""),
            "erp_updated": 1,
            "details": f"Match: {m.get('method', '')} | EDD: {m.get('edd', '')}",
        })


# ==========================================
# EMAIL
# ==========================================
def send_email(doc, excel_path, sarom_data=None):
    """Email the report using smtplib."""
    if sarom_data is None:
        sarom_data = {"matched": [], "total": 0, "in_stock": 0, "out_of_stock": 0, "low_stock": 0}
    log("  Sending email...")
    try:
        import smtplib
        from email.mime.multipart import MIMEMultipart
        from email.mime.text import MIMEText
        from email.mime.base import MIMEBase
        from email import encoders

        email_account = frappe.get_doc("Email Account", "ZIPCushions Support")
        smtp_server = email_account.smtp_server
        smtp_port = int(email_account.smtp_port or 587)
        login_id = email_account.login_id or email_account.email_id
        password = email_account.get_password(fieldname="password")
        sender = email_account.email_id

        subject = f"Vendor Stock Report - {doc.run_date} [{doc.overall_status}]"
        body = f"""<html><body>
        <h3>Vendor Stock Scrape Report</h3>
        <p><b>Date:</b> {doc.run_date} | <b>Status:</b> {doc.overall_status} | <b>Duration:</b> {doc.run_duration_mins} mins</p>
        <table border="1" cellpadding="5" cellspacing="0" style="border-collapse:collapse">
            <tr style="background:#2F5496;color:white"><th>Vendor</th><th>Status</th><th>Total</th><th>In Stock</th><th>Out</th><th>Low</th></tr>
            <tr><td>Agora</td><td>{doc.agora_status}</td><td>{doc.agora_total}</td><td>{doc.agora_in_stock}</td><td>{doc.agora_out_of_stock}</td><td>{doc.agora_low_stock}</td></tr>
            <tr><td>Linen Craft</td><td>{doc.lc_status}</td><td>{doc.lc_total}</td><td>{doc.lc_in_stock}</td><td>{doc.lc_out_of_stock}</td><td>{doc.lc_low_stock}</td></tr>
            <tr><td>DDécor</td><td>{doc.ddecor_status}</td><td>{doc.ddecor_total}</td><td>{doc.ddecor_in_stock}</td><td>{doc.ddecor_out_of_stock}</td><td>{doc.ddecor_low_stock}</td></tr>
            <tr><td>Sarom</td><td>{"Success" if sarom_data.get("matched") else "Failed"}</td><td>{sarom_data.get("total", 0)}</td><td>{sarom_data.get("in_stock", 0)}</td><td>{sarom_data.get("out_of_stock", 0)}</td><td>{sarom_data.get("low_stock", 0)}</td></tr>
        </table>
        <p><b>ERP Items Updated:</b> {doc.erp_items_updated}</p>
        <p>Full log: <a href="https://erp.cozycornerpatios.com/app/vendor-stock-scrape-log/{doc.name}">{doc.name}</a></p>
        </body></html>"""

        msg = MIMEMultipart()
        msg["From"] = sender
        msg["To"] = ", ".join(EMAIL_RECIPIENTS)
        msg["Subject"] = subject
        msg.attach(MIMEText(body, "html"))

        if excel_path and os.path.exists(excel_path):
            with open(excel_path, "rb") as f:
                part = MIMEBase("application", "octet-stream")
                part.set_payload(f.read())
                encoders.encode_base64(part)
                part.add_header("Content-Disposition", f"attachment; filename={os.path.basename(excel_path)}")
                msg.attach(part)

        server = smtplib.SMTP(smtp_server, smtp_port)
        server.starttls()
        server.login(login_id, password)
        server.sendmail(sender, EMAIL_RECIPIENTS, msg.as_string())
        server.quit()

        doc.email_sent = 1
        doc.email_recipients = ", ".join(EMAIL_RECIPIENTS)
        doc.save(ignore_permissions=True)
        frappe.db.commit()
        log(f"  Email sent to {', '.join(EMAIL_RECIPIENTS)}")

    except Exception as e:
        doc.email_sent = 0
        doc.email_error = str(e)[:500]
        doc.save(ignore_permissions=True)
        frappe.db.commit()
        log(f"  Email failed: {e}")


# ==========================================
# MAIN
# ==========================================
def main(triggered_by="Scheduled"):
    log(f"\n{'=' * 60}")
    log(f"VENDOR STOCK SCRAPE - {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    log(f"{'=' * 60}\n")

    errors = []
    agora_data = {"products": [], "total": 0}
    lc_data = {"products": [], "total": 0}
    ddecor_data = {"results": [], "total": 0}
    sarom_data = {"matched": [], "unmatched": [], "total": 0}
    erp_updated = 0
    excel_path = None

    # Create log doc
    scrape_log = create_scrape_log(triggered_by)

    try:
        # 1. AGORA
        log("=" * 50)
        log("1. AGORA")
        log("=" * 50)
        try:
            agora_data = agora.scrape(log_fn=log)
            erp_updated += sync_agora_to_erp(agora_data["products"])
        except Exception as e:
            log(f"  AGORA FAILED: {e}")
            errors.append(f"Agora: {str(e)[:200]}")
            agora_data["error"] = str(e)[:500]

        # 2. LINEN CRAFT
        log(f"\n{'=' * 50}")
        log("2. LINEN CRAFT")
        log("=" * 50)
        try:
            lc_data = linen_craft.scrape(log_fn=log)
            erp_updated += sync_linen_craft_to_erp(lc_data["products"])
        except Exception as e:
            log(f"  LINEN CRAFT FAILED: {e}")
            errors.append(f"Linen Craft: {str(e)[:200]}")
            lc_data["error"] = str(e)[:500]

        # 3. DDECOR
        log(f"\n{'=' * 50}")
        log("3. DDECOR")
        log("=" * 50)
        try:
            ddecor_data = ddecor.scrape(log_fn=log)
            erp_updated += sync_ddecor_to_erp(ddecor_data["results"])
        except Exception as e:
            log(f"  DDECOR FAILED: {e}")
            errors.append(f"DDécor: {str(e)[:200]}")
            ddecor_data["error"] = str(e)[:500]

        # 4. SAROM
        log(f"\n{'=' * 50}")
        log("4. SAROM")
        log("=" * 50)
        try:
            sarom_data = sarom.scrape(log_fn=log)
            if sarom_data.get("matched"):
                erp_updated += sync_sarom_to_erp(sarom_data["matched"])
        except Exception as e:
            log(f"  SAROM FAILED: {e}")
            errors.append(f"Sarom: {str(e)[:200]}")
            sarom_data["error"] = str(e)[:500]

        # 5. EXCEL
        log(f"\n{'=' * 50}")
        log("5. EXCEL REPORT")
        log("=" * 50)
        try:
            excel_path = generate_excel(agora_data, lc_data, ddecor_data, sarom_data)
        except Exception as e:
            log(f"  EXCEL FAILED: {e}")
            errors.append(f"Excel: {str(e)[:200]}")

        # 6. UPDATE LOG
        log(f"\n{'=' * 50}")
        log("6. UPDATING SCRAPE LOG")
        log("=" * 50)
        update_scrape_log(scrape_log, agora_data, lc_data, ddecor_data, sarom_data, erp_updated, excel_path, errors)

        # 7. EMAIL
        log(f"\n{'=' * 50}")
        log("7. EMAIL")
        log("=" * 50)
        send_email(scrape_log, excel_path, sarom_data)

    except Exception as e:
        log(f"\nFATAL ERROR: {e}")
        log(traceback.format_exc())
        try:
            scrape_log.overall_status = "Failed"
            scrape_log.error_summary = str(e)[:500]
            scrape_log.run_end = frappe.utils.now()
            scrape_log.save(ignore_permissions=True)
            frappe.db.commit()
        except:
            pass

    log(f"\n{'=' * 60}")
    log(f"DONE - Status: {scrape_log.overall_status}")
    log(f"Log: {scrape_log.name}")
    log(f"ERP items updated: {erp_updated}")
    log(f"{'=' * 60}")


if __name__ == "__main__":
    triggered = sys.argv[1] if len(sys.argv) > 1 else "Scheduled"
    main(triggered_by=triggered)
